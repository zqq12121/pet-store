"""小批量知识导入：先校验完整结构，再把合法行与任务结果一次提交。"""
import asyncio
import csv
import hashlib
import io
import json
import zipfile
from typing import Literal
from uuid import UUID

from fastapi import Depends, Header, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from .business import ApiError
from .store import now, uid

HEADERS = ['title', 'category', 'petType', 'breedNames', 'format', 'content', 'question', 'answer', 'sourceName', 'sourceUrl']
EXAMPLES = [
    ['布偶猫日常梳毛', 'grooming', 'cat', '布偶猫', 'article', '轻柔梳理，发现皮肤异常时建议就医。', '', '', '门店护理手册', ''],
    ['幼犬基础训练', 'training', 'dog', '', 'qa', '', '幼犬如何学习坐下？', '短时练习并使用奖励引导。', '历史问答，经店主审核', ''],
]
MIMES = {'csv': 'text/csv', 'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}
LIMIT = 10 * 1024 * 1024


def template(format):
    if format == 'csv':
        output = io.StringIO()
        csv.writer(output).writerows([HEADERS, *EXAMPLES])
        return output.getvalue().encode('utf-8-sig')
    book = Workbook()
    sheet = book.active
    sheet.title = 'knowledge'
    for row in [HEADERS, *EXAMPLES]:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def parse_rows(data, format):
    book = None
    try:
        if format == 'csv':
            rows = csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''), strict=True)
        else:
            # 限制解压体积，拒绝宏和外部链接；不执行或读取公式的缓存结果。
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(item.file_size for item in entries) > 40 * 1024 * 1024:
                    raise ApiError(400, 'IMPORT_FILE_INVALID', 'Excel 解压后过大')
                if any('vbaproject' in item.filename.lower() or 'externallinks/' in item.filename.lower() for item in entries):
                    raise ApiError(400, 'IMPORT_FILE_INVALID', 'Excel 不能包含宏或外部链接')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            if book.sheetnames != ['knowledge']:
                raise ApiError(400, 'IMPORT_FILE_INVALID', 'Excel 只能包含 knowledge 工作表')
            sheet = book['knowledge']
            if (sheet.max_row or 0) > 5001 or (sheet.max_column or 0) > len(HEADERS):
                raise ApiError(400, 'IMPORT_ROW_LIMIT_EXCEEDED', 'Excel 范围超过 5000 行或模板列数')
            rows = ([cell.value for cell in row] for row in sheet.iter_rows())
        if next(rows, None) != HEADERS:
            raise ApiError(400, 'IMPORT_INVALID_HEADER', '列头与模板不一致，请下载最新模板')
        result = []
        for number, row in enumerate(rows, 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            result.append((number, row))
            if len(result) > 5000:
                raise ApiError(400, 'IMPORT_ROW_LIMIT_EXCEEDED', '最多导入 5000 条非空数据行')
        if not result:
            raise ApiError(400, 'IMPORT_FILE_INVALID', '文件没有可导入的数据行')
        return result
    except ApiError:
        raise
    except Exception as exc:
        raise ApiError(400, 'IMPORT_FILE_INVALID', '文件损坏或格式不正确，请使用 UTF-8 CSV 或模板 Excel') from exc
    finally:
        if book:
            book.close()


class Imports:
    def __init__(self, repository, input_model):
        self.repository, self.input_model = repository, input_model
        self.tasks = set()
        with repository.store.db() as db:
            db.execute('''CREATE TABLE IF NOT EXISTS knowledge_imports (
                id TEXT PRIMARY KEY, actor TEXT NOT NULL, key TEXT NOT NULL, digest TEXT NOT NULL,
                data TEXT NOT NULL, UNIQUE(actor,key))''')
            # 进程中断时整个批次事务会回滚；不把未完成导入显示为成功。
            for row in db.execute('SELECT id,data FROM knowledge_imports').fetchall():
                job = json.loads(row['data'])
                if job['status'] in ('queued', 'running'):
                    job.update(status='failed', errorCode='IMPORT_INTERRUPTED', errorMessage='服务重启，导入已中断，请重新上传', finishedAt=now())
                    db.execute('UPDATE knowledge_imports SET data=? WHERE id=?', (json.dumps(job), row['id']))

    def get(self, job_id):
        with self.repository.store.db() as db:
            row = db.execute('SELECT data FROM knowledge_imports WHERE id=?', (job_id,)).fetchone()
        if not row:
            raise ApiError(404, 'KNOWLEDGE_JOB_NOT_FOUND', '导入任务不存在')
        return json.loads(row[0])

    def list(self):
        with self.repository.store.db() as db:
            return [json.loads(row[0]) for row in db.execute('SELECT data FROM knowledge_imports ORDER BY rowid DESC LIMIT 100')]

    def submit(self, actor, key, data, format):
        digest = hashlib.sha256(format.encode()+data).hexdigest()
        with self.repository.store.db() as db:
            db.execute('BEGIN IMMEDIATE')
            old = db.execute('SELECT digest,data FROM knowledge_imports WHERE actor=? AND key=?', (actor, key)).fetchone()
            if old:
                if old['digest'] != digest:
                    raise ApiError(409, 'IDEMPOTENCY_CONFLICT', '同一个导入编号不能用于不同文件')
                return json.loads(old['data'])
            if len(self.tasks) >= 2:
                raise ApiError(429, 'RATE_LIMITED', '导入任务较多，请稍后重试')
            job = dict(id=uid('kjob'), type='import', scope='knowledge', status='queued', totalCount=None,
                       processedCount=0, successCount=0, failedCount=0, errors=[], createdAt=now(),
                       startedAt=None, finishedAt=None, errorCode=None, errorMessage=None, createdBy=actor)
            db.execute('INSERT INTO knowledge_imports VALUES (?,?,?,?,?)', (job['id'], actor, key, digest, json.dumps(job)))
        task = asyncio.create_task(self.run(job, actor, data, format))
        self.tasks.add(task)
        task.add_done_callback(self.tasks.discard)
        return job

    async def stop(self):
        # 等待有界解析及事务落库完成，避免停止过程中产生半批结果。
        await asyncio.gather(*self.tasks, return_exceptions=True)

    async def run(self, job, actor, data, format):
        job.update(status='running', startedAt=now())
        self.save_job(job)
        try:
            rows = await asyncio.to_thread(parse_rows, data, format)
            valid, errors = [], []
            for number, row in rows:
                try:
                    if len(row) != len(HEADERS) or any(value is not None and not isinstance(value, str) for value in row):
                        raise ValueError('每行必须与模板列数一致，所有单元格须为文本')
                    if any((value or '').lstrip().startswith(('=', '+', '-', '@')) for value in row):
                        raise ValueError('单元格不能以公式字符开头')
                    body = {key: (value.strip() if value else None) for key, value in zip(HEADERS, row)}
                    body['breedNames'] = [name.strip() for name in (body['breedNames'] or '').split(';') if name.strip()]
                    valid.append(self.input_model.model_validate(body).model_dump())
                except (ValueError, ValidationError) as exc:
                    # 校验错误只返回字段和提示，不回传整行文本或 URL 中的敏感内容。
                    issue = exc.errors(include_input=False, include_context=False)[0] if isinstance(exc, ValidationError) else None
                    field = str(issue['loc'][0]) if issue and issue['loc'] else None
                    # 自定义校验提示均由 KnowledgeInput 固定定义，可说明缺失问答等具体原因。
                    message = issue['msg'].removeprefix('Value error, ') if issue and issue['type'] == 'value_error' else (
                        '字段不符合模板要求，请核对该字段' if issue else str(exc))
                    errors.append(dict(rowNumber=number, field=field, code='VALIDATION_ERROR', message=message))
            job.update(totalCount=len(rows), processedCount=len(rows), successCount=len(valid), failedCount=len(errors), errors=errors,
                       status='partial_succeeded' if valid and errors else 'succeeded' if valid else 'failed', finishedAt=now())
            # 合法草稿与终态共用一个事务；重放任务绝不重复新增知识。
            with self.repository.store.db() as db:
                db.execute('BEGIN IMMEDIATE')
                for body in valid:
                    self.repository.save(body, actor, db=db)
                db.execute('UPDATE knowledge_imports SET data=? WHERE id=?', (json.dumps(job), job['id']))
        except Exception as exc:
            job.update(status='failed', totalCount=None, processedCount=0, successCount=0, failedCount=0, errors=[], finishedAt=now(),
                       errorCode=exc.code if isinstance(exc, ApiError) else 'IMPORT_FAILED',
                       errorMessage=exc.message if isinstance(exc, ApiError) else '导入失败，请稍后重新上传')
            self.save_job(job)

    def save_job(self, job):
        with self.repository.store.db() as db:
            db.execute('UPDATE knowledge_imports SET data=? WHERE id=?', (json.dumps(job), job['id']))


def register_import_api(app, imports, administrator, envelope, page):
    @app.get('/api/v1/admin/knowledge/import-template')
    async def download(format: Literal['csv', 'xlsx'] = Query(...), actor=Depends(administrator)):
        return Response(template(format), media_type=MIMES[format],
                        headers={'Content-Disposition': f'attachment; filename="knowledge-import-template.{format}"'})

    @app.post('/api/v1/admin/knowledge/import-jobs', status_code=202)
    async def submit(request: Request, response: Response, format: Literal['csv', 'xlsx'] = Query(...),
                     key: UUID = Header(..., alias='Idempotency-Key'), actor=Depends(administrator)):
        if request.headers.get('content-type', '').split(';')[0] != MIMES[format]:
            raise ApiError(415, 'UNSUPPORTED_FILE_TYPE', '请上传与格式一致的 CSV 或 Excel 文件')
        data = bytearray()
        # 分块限制请求体，避免先把超大文件完整读入内存；不接收路径或下载地址。
        async for chunk in request.stream():
            if len(data)+len(chunk) > LIMIT:
                raise ApiError(413, 'FILE_TOO_LARGE', '文件不能超过 10 MB')
            data.extend(chunk)
        job = imports.submit(actor.owner, str(key), bytes(data), format)
        response.headers['Location'] = '/api/v1/admin/knowledge/jobs/'+job['id']
        return envelope({k: v for k, v in job.items() if k != 'errors'})

    @app.get('/api/v1/admin/knowledge/jobs')
    async def jobs(actor=Depends(administrator)):
        return envelope([{k: v for k, v in job.items() if k != 'errors'} for job in imports.list()])

    @app.get('/api/v1/admin/knowledge/jobs/{job_id}')
    async def job(job_id: str, actor=Depends(administrator)):
        return envelope({k: v for k, v in imports.get(job_id).items() if k != 'errors'})

    @app.get('/api/v1/admin/knowledge/jobs/{job_id}/errors')
    async def errors(job_id: str, number: int = Query(1, alias='page', ge=1), pageSize: int = Query(20, ge=1, le=100),
                     actor=Depends(administrator)):
        return page(imports.get(job_id)['errors'], number, pageSize)
